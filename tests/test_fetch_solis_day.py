"""fetch_solis_day: Solis exports + PEA Log → processed CSVs."""

from __future__ import annotations

import csv
import os
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from momsolar import fetch_solis_day as f
from momsolar.schema import DAILY_COLUMNS, FIVE_MIN_COLUMNS

# The header of the real All-History export, trimmed to the columns we read + a Chinese one.
HIST_HEADER = [
    "Number",
    "Time",
    "Working State",
    "Alarm Code",
    "DC Voltage PvMPPT1(V)",
    "DC Power PvMPPT1(W)",
    "DC Voltage PvMPPT2(V)",
    "DC Power PvMPPT2(W)",
    "Today Yield(kWh)",
    "Inverter Internal Operating Ambient Temperature(℃)",
    "Grid Total Active Power(W)",
    "Daily Energy from Grid(kWh)",
    "Battery Power(W)",
    "Today Energy to Battery(kWh)",
    "Today Energy from Battery(kWh)",
    "Battery SOC(%)",
    "Battery SOH(%)",
    "Today Grid Load Consumption(kWh)",
    "Total Grid Load Consumption(kWh)",
    "Grid load power(W)",
    "Today Backup Load Consumption(kWh)",
    "Total Backup Load Consumption(kWh)",
    "Backup Active Power L1(W)",
    "Smart Port active power L1(W)",
    "Gen active power L1(W)",
    "Smart Port AC Coupled Input Power(W)",
    "外部逆变器功率 L1(W)",
]


def _hist_row(n, t, mppt1, mppt2, bat, grid, soc, yld="0", day="23/09/2026"):
    return [
        str(n),
        f"{day} {t} (UTC+07:00)",
        "Normal",
        " ",
        "260.3",
        mppt1,
        "257.0",
        mppt2,
        yld,
        "56.2",
        grid,
        "0.100",
        bat,
        "10.700",
        "4.600",
        soc,
        "99.000",
        "0.100",
        "213.000",
        "0",
        "7.200",
        "3452.000",
        "1370.000",
        "0",
        "0",
        "0",
        "0",
    ]


def write_history_xlsx(path: Path, rows: list[list[str]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inverter All History Report__23"
    ws.append(["Inverter All History Report_23/09/2026_Report"])
    ws.append([])
    ws.append(["Inverter SN:1234567890123456", None, "Inverter type:3197"])
    ws.append(["Collector SN:ABCDEF0123456789", None, "Belongs to plant:Demo"])
    ws.append(["Rated Power:10000Wp"])
    ws.append([])
    ws.append([])
    ws.append(HIST_HEADER)
    for r in rows:
        ws.append(r)
    # Second sheet: power only — must be skipped.
    ws2 = wb.create_sheet("Inverter History Report__23-09-")
    ws2.append(["Number", "Time", "Working State", "Alarm Code", "Total Inverter Power(W)"])
    wb.save(path)
    return path


def write_monthly_xlsx(path: Path, rows: list[list]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Inverter_Report_2026-09_Report"])
    ws.append([])
    ws.append(["Inverter SN:1234567890123456"])
    ws.append([])
    ws.append(["Full Load Hours(h):43.25"])
    ws.append([])
    ws.append(
        [
            "Inverter SN",
            "Number",
            "Time",
            "Today Yield(kWh)",
            "Today Full Load Hours(h)",
            "Energy to Grid(kWh)",
            "Energy from Grid(kWh)",
            "Energy to Battery(kWh)",
            "Energy from Battery(kWh)",
            "Load Consumption(kWh)",
            "Generation(kWh)",
            "GEN(kWh)",
            "Smart Load(kWh)",
            "AC Coupled(kWh)",
        ]
    )
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def read(path: Path) -> list[dict]:
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


# ── parsing helpers ─────────────────────────────────────────────────────────
def test_parse_solis_time_strips_zone():
    assert f.parse_solis_time("23/09/2026 00:03:36 (UTC+07:00)") == datetime(2026, 9, 23, 0, 3, 36)


# ── 5-min ───────────────────────────────────────────────────────────────────
def test_five_min_rows_maps_columns_and_sums_pv(tmp_path):
    p = write_history_xlsx(
        tmp_path / "Inverter_Report_1234567890123456_23-09-2026_20260923132227.xlsx",
        [
            _hist_row(1, "00:03:36", "0", "0", "-410.000", "0", "62"),
            _hist_row(2, "12:18:49", "729", "720", "0", "-20.000", "100", yld="14.600"),
        ],
    )
    rows = f.five_min_rows(p)
    assert len(rows) == 2
    last = rows[-1]
    assert last["Time"] == "2026-09-23 12:18:49"
    assert last["PV(W)"] == "1449"
    assert last["Grid(W)"] == "-20"
    assert last["SOC(%)"] == "100"
    assert last["Today Yield(kWh)"] == "14.6"
    assert last["Today Backup Load(kWh)"] == "7.2"
    assert last["Total Grid Load(kWh)"] == "213"
    assert rows[0]["Battery(W)"] == "-410"
    assert rows[0]["Alarm Code"] == ""  # Solis writes a single space when there is no alarm
    assert set(rows[0]) == set(FIVE_MIN_COLUMNS)


def test_five_min_rows_rejects_other_files(tmp_path):
    p = write_monthly_xlsx(tmp_path / "x.xlsx", [])
    with pytest.raises(ValueError):
        f.five_min_rows(p)


def test_classify_by_header_not_name(tmp_path):
    hist = write_history_xlsx(
        tmp_path / "a.xlsx", [_hist_row(1, "00:03:36", "0", "0", "0", "0", "62")]
    )
    monthly = write_monthly_xlsx(tmp_path / "Inverter History Report_1.xlsx", [])
    assert f.classify(f.read_sheets(hist))[0] == "5min"
    assert f.classify(f.read_sheets(monthly))[0] == "daily"
    old = {"s": [["Number", "Time", "Battery SOC(%)", "Backup Load Power(W)"]]}
    assert f.classify(old)[0] == "5min-old"
    assert f.classify({"s": [["Plant ID", "Plant Name"]]}) is None


def test_multi_day_export_and_duplicate_readings(tmp_path):
    p = write_history_xlsx(
        tmp_path / "multi.xlsx",
        [
            _hist_row(1, "23:58:40", "0", "0", "0", "0", "40", day="22/09/2026"),
            _hist_row(2, "00:03:36", "0", "0", "0", "0", "39"),
            _hist_row(3, "00:03:36", "0", "0", "0", "0", "39"),  # Solis repeats a reading
        ],
    )
    rows = f.five_min_rows(p)
    assert [r["Time"] for r in rows] == ["2026-09-22 23:58:40", "2026-09-23 00:03:36"]


def test_newer_export_of_a_day_wins(tmp_path):
    raw, out = tmp_path / "raw", tmp_path / "out"
    raw.mkdir()
    a = write_history_xlsx(raw / "a.xlsx", [_hist_row(1, "00:03:36", "0", "0", "0", "0", "62")])
    b = write_history_xlsx(
        raw / "b.xlsx",
        [
            _hist_row(1, "00:03:36", "0", "0", "0", "0", "62"),
            _hist_row(2, "00:08:36", "0", "0", "0", "0", "61"),
        ],
    )
    os.utime(a, (1, 1))  # a is older
    f.run(raw, out, log=lambda *_: None)
    assert len(read(out / "momhome" / "5min.csv")) == 2
    os.utime(b, (0, 0))  # now a (one reading) is the newest export of that day
    f.run(raw, out, log=lambda *_: None)
    assert len(read(out / "momhome" / "5min.csv")) == 1


def test_replace_days_swaps_whole_day():
    existing = [{"Time": "2026-09-22 23:58:00"}, {"Time": "2026-09-23 00:03:36"}]
    new = [{"Time": "2026-09-23 00:04:10"}]
    out = f.replace_days(existing, new)
    assert [r["Time"] for r in out] == ["2026-09-22 23:58:00", "2026-09-23 00:04:10"]


# ── daily ───────────────────────────────────────────────────────────────────
def test_daily_rows_iso_dates(tmp_path):
    p = write_monthly_xlsx(
        tmp_path / "Inverter_Report_2026-09_Report.xlsx",
        [["SN", "23", "23/09/2026", 14.1, 1.41, 0.0, 0.11, 10.0, 4.0, 7.0, 14.1, 0, 0, 0]],
    )
    (row,) = f.daily_rows(p)
    assert row["Time"] == "2026-09-23"
    assert row["Energy from Grid(kWh)"] == "0.11"
    assert row["Load Consumption(kWh)"] == "7"
    assert set(row) == set(DAILY_COLUMNS)


# ── end to end ──────────────────────────────────────────────────────────────
def test_run_writes_and_merges(tmp_path):
    raw, out = tmp_path / "raw", tmp_path / "out"
    raw.mkdir()
    write_history_xlsx(
        raw / "Inverter_Report_1234567890123456_23-09-2026_20260923132227.xlsx",
        [_hist_row(1, "00:03:36", "0", "0", "-410.000", "0", "62")],
    )
    write_monthly_xlsx(
        raw / "Inverter_Report_2026-09_Report.xlsx",
        [["SN", "1", "01/09/2026", 17.0, 1.7, 0.0, 0.27, 8.0, 9.0, 15.0, 17.0, 0, 0, 0]],
    )
    pea = tmp_path / "pea.csv"
    pea.write_text("Month Year,หน่วย,จำนวนเงิน\n31/8/2026,38,165.03\n", encoding="utf-8")

    counts = f.run(raw, out, str(pea), log=lambda *_: None)
    assert counts == {"momhome/5min.csv": 1, "momhome/daily.csv": 1, "momhome/bills.csv": 1}
    assert read(out / "momhome" / "bills.csv")[0]["Usage Month"] == "8"

    # Second run with a new day keeps the first day's rows.
    write_monthly_xlsx(
        raw / "Inverter_Report_2026-09_Report.xlsx",
        [["SN", "2", "02/09/2026", 18.6, 1.86, 0.0, 0.39, 10.0, 7.0, 14.0, 18.6, 0, 0, 0]],
    )
    f.run(raw, out, log=lambda *_: None)
    assert [r["Time"] for r in read(out / "momhome" / "daily.csv")] == ["2026-09-01", "2026-09-02"]


def test_home_option_picks_the_folder(tmp_path):
    raw, out = tmp_path / "raw", tmp_path / "out"
    raw.mkdir()
    write_monthly_xlsx(
        raw / "m.xlsx", [["SN", "1", "01/09/2026", 17.0, 1.7, 0, 0.27, 8, 9, 15, 17, 0, 0, 0]]
    )
    assert f.run(raw, out, log=lambda *_: None, home="other") == {"other/daily.csv": 1}


def test_main_requires_an_input(capsys):
    with pytest.raises(SystemExit):
        f.main([])
